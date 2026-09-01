"""文档解析器：把上传的文件统一转成纯文本。

支持 PDF / Markdown / TXT / Word(docx) / Excel(xlsx)，统一入口按扩展名分发。
"""
import io

from pypdf import PdfReader


def parse_pdf(content: bytes) -> str:
    """解析 PDF：逐页提取文字后拼接。

    Args:
        content: 文件的二进制内容。
    Returns:
        全文纯文本，页与页之间用换行分隔。
    """
    reader = PdfReader(io.BytesIO(content))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text.strip())
    return "\n".join(pages)


def parse_text(content: bytes) -> str:
    """解析纯文本类文件（Markdown / txt）：本身就是文本，直接解码。"""
    return content.decode("utf-8", errors="ignore")


def parse_docx(content: bytes) -> str:
    """解析 Word 文档：逐段落提取文字（python-docx）。"""
    from docx import Document

    doc = Document(io.BytesIO(content))
    return "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())


def parse_xlsx(content: bytes) -> str:
    """解析 Excel 表格：每个工作表按行输出，单元格用制表符分隔（openpyxl）。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"# 工作表: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            vals = [str(c).strip() for c in row if c is not None]
            if vals:
                lines.append("\t".join(vals))
    return "\n".join(lines)


_PARSERS = {
    ".pdf": parse_pdf,
    ".md": parse_text,
    ".txt": parse_text,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
}


def parse_file(filename: str, content: bytes) -> str:
    """统一入口：根据扩展名分发到对应的解析器。

    Raises:
        ValueError: 不支持的文件类型。
    """
    ext = "." + filename.rsplit(".", 1)[-1].lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"不支持的文件类型: {ext}，当前支持 {list(_PARSERS)}")
    return parser(content)
